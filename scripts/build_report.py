"""Build the Chinese project report from Markdown using python-docx.

Run with the Codex bundled Python runtime. This script does not read private
trading rows or credentials. Optional audit input contains aggregate counts only.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "標楷體"
LABELS = {
    "NO_OBSERVED_POSITION": "可見初始單",
    "ADD_ADVERSE": "價格不利時同向加單",
    "ADD_FAVORABLE": "價格有利時同向加單",
    "ADD_FLAT": "同價加單",
    "OPPOSITE": "反向持倉",
    "MIXED": "混合持倉",
    "TIME_TIE": "時間順序不明",
}


def set_font_properties(properties: Any, size: int = 12) -> None:
    """Set all Word font slots and remove inherited theme font/color overrides."""
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), FONT)
    for tag in ("sz", "szCs"):
        node = properties.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            properties.append(node)
        node.set(qn("w:val"), str(size * 2))
    color = properties.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        properties.append(color)
    color.attrib.clear()
    color.set(qn("w:val"), "000000")


def format_run(run: Any, size: int = 12, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = bold
    set_font_properties(run._element.get_or_add_rPr(), size)


def create_document() -> Any:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    section.footer_distance = Inches(0.3)
    for style in doc.styles:
        set_font_properties(style.element.get_or_add_rPr(), 12)
        if style.type == 1:
            style.paragraph_format.widow_control = True
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is not None:
                set_font_properties(rpr)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13)):
        style = doc.styles[name]
        set_font_properties(style.element.get_or_add_rPr(), size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    title = doc.styles["Title"]
    title.paragraph_format.line_spacing = 1.1
    for element in list(title.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    set_font_properties(rpr)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    doc.core_properties.title = "結合歷史交易案例檢索與對話式風險提示的黃金交易輔助平台"
    doc.core_properties.subject = "專題報告與階段擴充規劃"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    return doc


def add_paragraph(doc: Any, text: str, style: str | None = None) -> Any:
    paragraph = doc.add_paragraph(style=style)
    size = 18 if style == "Title" else 15 if style == "Heading 1" else 13 if style == "Heading 2" else 12
    format_run(paragraph.add_run(text), size, bool(style))
    return paragraph


def add_table(doc: Any, rows: list[list[str]]) -> None:
    columns = len(rows[0])
    table = doc.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    ratios = {2: [0.70, 0.30], 3: [0.22, 0.37, 0.41], 4: [0.34, 0.22, 0.22, 0.22]}
    widths = [Inches(7.1 * n) for n in ratios.get(columns, [1 / columns] * columns)]
    for col, width in zip(table.columns, widths):
        col.width = width
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn(f"w:{key}"), value)
        borders.append(node)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "bottom", "left", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "80" if edge in ("top", "bottom") else "100")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)
    for index, values in enumerate(rows):
        row = table.add_row()
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(repeat)
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[col]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.keep_with_next = False
            numeric = bool(re.fullmatch(r"[\d,]+", value))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric else WD_ALIGN_PARAGRAPH.LEFT
            format_run(paragraph.add_run(value), bold=index == 0)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDEDED")
                cell._tc.get_or_add_tcPr().append(shade)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0


def math_run(text: str) -> Any:
    run = OxmlElement("m:r")
    props = OxmlElement("w:rPr")
    set_font_properties(props)
    run.append(props)
    node = OxmlElement("m:t")
    node.text = text
    run.append(node)
    return run


def fraction(numerator: str, denominator: str) -> Any:
    node = OxmlElement("m:f")
    for tag, text in (("num", numerator), ("den", denominator)):
        part = OxmlElement(f"m:{tag}")
        part.append(math_run(text))
        node.append(part)
    return node


def add_equation(doc: Any, kind: str) -> None:
    """Editable native Word math with deliberate linear or fraction structure."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement("m:oMath")
    if kind == "average":
        math.append(math_run("同向平均進場價 = "))
        math.append(fraction("各筆進場價 × 手數的加總", "同向總手數"))
    elif kind == "pnl":
        math.append(math_run("多單情境損益 = 各筆（情境賣出價 − 進場價）× 手數 × 每手盎司數的加總"))
    elif kind == "stop":
        math.append(math_run("新增單停損金額 = |進場價 − 停損價| × 新增手數 × 每手盎司數"))
    elif kind == "rr":
        math.append(math_run("計畫風險報酬比 = "))
        math.append(fraction("|停利價 − 進場價|", "|進場價 − 停損價|"))
    else:
        raise ValueError(f"Unsupported equation: {kind}")
    paragraph._p.append(math)


def update_audit_summary(markdown: str, summary_path: Path) -> str:
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    counts = summary["behavior_counts"]
    total = summary["gold_rows"]
    if sum(counts.values()) != total:
        raise ValueError("Audit behavior counts do not reconcile with gold_rows")
    if not summary["combined_counts_are_source_observations_not_unique_trades"]:
        raise ValueError("Report requires explicit observation-row semantics")
    formats = summary["gold_by_format"]
    first = summary["earliest_gold_open"][:10]
    last = summary["latest_gold_open"][:10]
    content = (
        f"2026 年 10 月 1 日刷新稽核涵蓋 {summary['source_files']:,} 個來源檔案：原有 "
        f"{summary['source_changes'].get('UNCHANGED', 0):,} 個檔案內容均未變，另新增 "
        f"{summary['source_changes'].get('ADDED', 0):,} 個 CSV。黃金可見來源觀察列共 {total:,} 列，"
        f"HTML 為 {formats['HTML']['observation_rows']:,} 列、Excel 為 {formats['XLSX']['observation_rows']:,} 列、"
        f"CSV 為 {formats['CSV']['observation_rows']:,} 列。可見黃金開倉時間範圍為 {first} 至 {last}。"
        "這些列含跨格式重複候選，不是去重後交易數，也不是完整歷史或獨立決策數。\n\n"
        "| 可觀察操作分類 | 來源觀察列數 |\n|---|---|\n"
    )
    content += "\n".join(f"| {label} | {counts.get(code, 0):,} |" for code, label in LABELS.items())
    content += f"\n| 合計 | {total:,} |\n\n"
    content += (
        f"跨來源相同簽章有 {summary['cross_source_signature_candidates']:,} 個候選，涉及 "
        f"{summary['cross_source_candidate_observation_rows']:,} 列；不能直接扣除候選數後宣稱獨立交易數。"
        f"{summary['html_declared_gap_files']} 份 HTML 有宣告列數與可見列數差距。CSV 中 "
        f"{formats['CSV']['behavior_counts'].get('TIME_TIE', 0):,} 列黃金操作順序不明，另有 "
        f"{summary['open_trade_snapshot_rows']} 列未平倉快照分開保存，不納入已實現交易統計。"
        "另有 PDF 交易表因方向不可讀，保留未分類紀錄，未反推方向或加入上述分群。"
        "原始檔案 SHA-256 刷新前後一致。"
    )
    pattern = r"(?s)(<!-- audit-summary:start -->).*?(<!-- audit-summary:end -->)"
    return re.sub(pattern, lambda match: f"{match[1]}\n{content}\n{match[2]}", markdown)


def build(markdown: str, output: Path) -> None:
    document = create_document()
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line.startswith("<!--"):
            continue
        if line.startswith(":::equation "):
            add_equation(document, line.split()[1])
            while index < len(lines) and lines[index].strip() != ":::":
                index += 1
            index += 1
            continue
        if line.startswith("|"):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index].strip())
                index += 1
            cells = [[item.strip() for item in row.strip("|").split("|")] for row in rows]
            cells = [row for row in cells if not all(re.fullmatch(r":?-+:?", item) for item in row)]
            add_table(document, cells)
            continue
        if line.startswith("### "):
            add_paragraph(document, line[4:], "Heading 2")
        elif line.startswith("## "):
            add_paragraph(document, line[3:], "Heading 1")
        elif line.startswith("# "):
            add_paragraph(document, line[2:], "Title")
        else:
            paragraph_lines = [line]
            while index < len(lines) and lines[index].strip() and not lines[index].startswith(("#", "|", ":::", "<!--")):
                paragraph_lines.append(lines[index].strip())
                index += 1
            add_paragraph(document, " ".join(paragraph_lines))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs" / "project-report.md")
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "黃金交易輔助平台專題報告_v1.docx")
    parser.add_argument("--audit-summary", type=Path, help="Optional refreshed aggregate summary JSON")
    args = parser.parse_args()
    markdown = args.source.read_text(encoding="utf-8")
    if args.audit_summary:
        markdown = update_audit_summary(markdown, args.audit_summary)
    build(markdown, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
